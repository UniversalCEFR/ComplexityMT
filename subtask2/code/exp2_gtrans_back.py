"""
Exp 2 - Google Translate BACKTRANSLATION classification pass.

Reads gtrans_templates/exp2_gtrans_back_input.xlsx (filled with text_back from
Google Sheets GOOGLETRANSLATE), runs UniversalCEFR on text_back, computes
    delta_l_forward = cefr_forward - cefr_orig
    delta_l_back    = cefr_back   - cefr_orig

Outputs (results/, slug=gtrans):
  exp2_{split}_back_{src}_to_{tgt}_gtrans.csv         per pair
  exp2_{split}_back_pairs_all_gtrans.csv              combined per split
  exp2_{split}_back_pair_summary_gtrans.{csv,json}    heatmap-ready aggregate

Matches the GPT-5.4 backtranslation schema so visualize_exp2_backtranslate.py
Just Works.
"""

import argparse
import json
import os
import time

import pandas as pd
import torch
from openpyxl import load_workbook
from transformers import AutoTokenizer, AutoModelForSequenceClassification

HF_TOKEN = os.environ.get("HF_TOKEN", "")
CLASSIFIER_MODEL = "UniversalCEFR/xlm-roberta-base-cefr-all-classifier"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(BASE_DIR, "gtrans_templates", "exp2_gtrans_back_input.xlsx")
RESULTS_DIR = os.path.join(BASE_DIR, "results")
os.makedirs(RESULTS_DIR, exist_ok=True)

SLUG = "gtrans"
CEFR_LEVELS = ["A1", "A2", "B1", "B2", "C1", "C2"]
CEFR_TO_INT = {l: i for i, l in enumerate(CEFR_LEVELS)}
INT_TO_CEFR = {i: l for i, l in enumerate(CEFR_LEVELS)}


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", default=DEFAULT_XLSX)
    p.add_argument("--split", nargs="+", default=["sentence", "document"],
                   choices=["sentence", "document"])
    p.add_argument("--hf-token", default=None)
    p.add_argument("--limit", type=int, default=None)
    p.add_argument("--forward-summary", default=None,
                   help="Path to exp2_{split}_pairs_all_gtrans.csv to merge in cefr_forward/delta_l_forward")
    return p.parse_args()


def load_xlsx_rows(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["backtranslations"]

    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), start=1):
        if row and row[0] == "row_id":
            header_row_idx = i
            break
    if header_row_idx is None:
        raise RuntimeError("Could not locate header row in %s" % xlsx_path)

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        rid, split, src, tgt, cefr, text_orig, text_fwd, text_back = row[:8]
        rows.append({
            "row_id": int(rid) if rid is not None and str(rid).strip() != "" else None,
            "split": split,
            "source_lang": src,
            "pivot_lang": tgt,
            "cefr_orig": cefr,
            "text_orig": text_orig,
            "text_forward": text_fwd,
            "text_back": text_back,
        })
    return rows


def load_classifier(hf_token=None):
    token = hf_token or HF_TOKEN
    os.environ["HUGGING_FACE_HUB_TOKEN"] = token
    os.environ["HF_TOKEN"] = token
    print("Loading classifier: %s" % CLASSIFIER_MODEL)
    tokenizer = AutoTokenizer.from_pretrained(CLASSIFIER_MODEL, use_auth_token=token)
    model = AutoModelForSequenceClassification.from_pretrained(CLASSIFIER_MODEL, use_auth_token=token)
    device = "cuda" if torch.cuda.is_available() else "cpu"
    model.to(device)
    model.eval()
    print("Classifier device: %s" % device)
    return tokenizer, model, device


def predict_cefr(tokenizer, model, device, text):
    if text is None or (isinstance(text, float) and pd.isna(text)) or str(text).strip() == "":
        return None
    inputs = tokenizer(str(text), return_tensors="pt", truncation=True, max_length=512)
    inputs = {k: v.to(device) for k, v in inputs.items()}
    with torch.no_grad():
        logits = model(**inputs).logits
    pred_idx = int(torch.argmax(logits, dim=-1).item())
    return model.config.id2label.get(pred_idx, INT_TO_CEFR.get(pred_idx, str(pred_idx)))


def classify_back(rows, tokenizer, model, device):
    total = len(rows)
    t0 = time.time()
    for i, r in enumerate(rows):
        r["cefr_back"] = predict_cefr(tokenizer, model, device, r["text_back"])
        if (i + 1) % 100 == 0 or (i + 1) == total:
            elapsed = time.time() - t0
            rate = (i + 1) / max(elapsed, 1e-6)
            eta = (total - (i + 1)) / max(rate, 1e-6)
            print("  classified %d/%d  (%.1f/s, elapsed %.1fs, ETA %.1fs)" % (
                i + 1, total, rate, elapsed, eta), flush=True)
    return rows


def attach_forward_predictions(rows, forward_summary_path):
    """Merge cefr_forward + delta_l_forward from the existing forward combined CSV."""
    if not forward_summary_path or not os.path.exists(forward_summary_path):
        print("  (no forward summary provided; cefr_forward/delta_l_forward left blank)")
        for r in rows:
            r["cefr_forward"] = None
            r["delta_l_forward"] = None
        return rows
    fwd = pd.read_csv(forward_summary_path)
    # Build lookup key: (source_lang, target_lang, text_orig)
    fwd["key"] = fwd["source_lang"] + "|" + fwd["target_lang"] + "|" + fwd["text_orig"].astype(str)
    lookup = dict(zip(fwd["key"], zip(fwd["cefr_pred"], fwd["delta_l"])))
    hit, miss = 0, 0
    for r in rows:
        key = "%s|%s|%s" % (r["source_lang"], r["pivot_lang"], r["text_orig"])
        if key in lookup:
            cefr_fwd, delta_fwd = lookup[key]
            r["cefr_forward"] = cefr_fwd
            r["delta_l_forward"] = delta_fwd
            hit += 1
        else:
            r["cefr_forward"] = None
            r["delta_l_forward"] = None
            miss += 1
    print("  forward merge: %d hits, %d misses" % (hit, miss))
    return rows


def compute_delta(rows):
    for r in rows:
        orig, back = r.get("cefr_orig"), r.get("cefr_back")
        r["delta_l_back"] = (CEFR_TO_INT[back] - CEFR_TO_INT[orig]
                             if orig in CEFR_TO_INT and back in CEFR_TO_INT else None)
    return rows


def save_per_pair(df_split, split):
    out_paths = []
    for (src, tgt), grp in df_split.groupby(["source_lang", "pivot_lang"]):
        out_path = os.path.join(RESULTS_DIR, "exp2_%s_back_%s_to_%s_%s.csv" % (split, src, tgt, SLUG))
        out = grp.copy()
        out["round_trip"] = "%s->%s->%s" % (src.upper(), tgt.upper(), src.upper())
        out["mt_model"] = "google_translate"
        out["format"] = split
        cols = ["round_trip", "source_lang", "pivot_lang", "format", "mt_model",
                "cefr_orig", "cefr_forward", "cefr_back",
                "text_orig", "text_forward", "text_back",
                "delta_l_forward", "delta_l_back"]
        out = out[[c for c in cols if c in out.columns]]
        out.to_csv(out_path, index=False, encoding="utf-8")
        out_paths.append(out_path)
    print("  wrote %d per-pair back CSVs for %s-level" % (len(out_paths), split))


def save_combined(df_split, split):
    out_path = os.path.join(RESULTS_DIR, "exp2_%s_back_pairs_all_%s.csv" % (split, SLUG))
    df = df_split.copy()
    df["round_trip"] = df["source_lang"].str.upper() + "->" + df["pivot_lang"].str.upper() + "->" + df["source_lang"].str.upper()
    df["mt_model"] = "google_translate"
    df["format"] = split
    cols = ["round_trip", "source_lang", "pivot_lang", "format", "mt_model",
            "cefr_orig", "cefr_forward", "cefr_back",
            "text_orig", "text_forward", "text_back",
            "delta_l_forward", "delta_l_back"]
    df = df[[c for c in cols if c in df.columns]]
    df.to_csv(out_path, index=False, encoding="utf-8")
    print("  wrote combined: %s (%d rows)" % (out_path, len(df)))
    return df


def write_pair_summary(df_all, split):
    valid = df_all.dropna(subset=["delta_l_back"]).copy()
    if valid.empty:
        return
    agg = (valid.groupby(["source_lang", "pivot_lang"])
                .agg(n=("delta_l_back", "count"),
                     mean_back=("delta_l_back", "mean"),
                     std_back=("delta_l_back", "std"),
                     median_back=("delta_l_back", "median"),
                     mean_forward=("delta_l_forward", "mean"))
                .reset_index())
    agg["pct_up_back"] = (valid.assign(u=(valid["delta_l_back"] > 0).astype(int))
                                .groupby(["source_lang", "pivot_lang"])["u"].mean().values)
    agg["pct_down_back"] = (valid.assign(d=(valid["delta_l_back"] < 0).astype(int))
                                  .groupby(["source_lang", "pivot_lang"])["d"].mean().values)
    agg["pct_preserved_back"] = (valid.assign(e=(valid["delta_l_back"] == 0).astype(int))
                                       .groupby(["source_lang", "pivot_lang"])["e"].mean().values)

    csv_path = os.path.join(RESULTS_DIR, "exp2_%s_back_pair_summary_%s.csv" % (split, SLUG))
    agg.to_csv(csv_path, index=False, encoding="utf-8")

    nested = {}
    for _, row in agg.iterrows():
        nested.setdefault(row["source_lang"], {})[row["pivot_lang"]] = {
            "n": int(row["n"]),
            "mean_delta_l_back": float(row["mean_back"]),
            "std_delta_l_back": float(row["std_back"]) if pd.notna(row["std_back"]) else None,
            "median_delta_l_back": float(row["median_back"]),
            "mean_delta_l_forward": float(row["mean_forward"]) if pd.notna(row["mean_forward"]) else None,
            "pct_up_back": float(row["pct_up_back"]),
            "pct_down_back": float(row["pct_down_back"]),
            "pct_preserved_back": float(row["pct_preserved_back"]),
        }
    json_path = os.path.join(RESULTS_DIR, "exp2_%s_back_pair_summary_%s.json" % (split, SLUG))
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(nested, f, indent=2)

    print("  wrote summary: %s" % csv_path)
    print("  wrote JSON:    %s" % json_path)


def main():
    args = parse_args()
    print("Reading %s" % args.xlsx)
    rows = load_xlsx_rows(args.xlsx)
    if args.limit:
        rows = rows[:args.limit]
    print("Loaded %d rows" % len(rows))

    tokenizer, model, device = load_classifier(hf_token=args.hf_token)
    rows = classify_back(rows, tokenizer, model, device)
    rows = compute_delta(rows)

    # Auto-detect forward summary per split
    for split in args.split:
        fwd_path = args.forward_summary or os.path.join(
            RESULTS_DIR, "exp2_%s_pairs_all_%s.csv" % (split, SLUG))
        split_rows = [r for r in rows if r["split"] == split]
        if not split_rows:
            print("No rows for split=%s" % split)
            continue
        split_rows = attach_forward_predictions(split_rows, fwd_path)
        df_split = pd.DataFrame(split_rows)
        save_per_pair(df_split, split)
        df_all = save_combined(df_split, split)
        write_pair_summary(df_all, split)

        valid = df_all.dropna(subset=["delta_l_back"])
        print("\n=== %s-level backtranslation summary ===" % split.upper())
        print("valid rows: %d / %d" % (len(valid), len(df_all)))
        if len(valid):
            print("overall mean delta_l_back:    %.3f" % valid["delta_l_back"].mean())
            if valid["delta_l_forward"].notna().any():
                print("overall mean delta_l_forward: %.3f"
                      % valid["delta_l_forward"].dropna().mean())


if __name__ == "__main__":
    main()
