"""
Exp 2 - Google Translate classification pass.

Reads the filled Google-Sheets workbook (gtrans_templates/exp2_gtrans_input.xlsx),
predicts CEFR on every text_translated cell with the UniversalCEFR classifier,
computes delta_l = cefr_pred - cefr_orig, and writes outputs in the same shape
as the GPT-5.4 pipeline but with slug='gtrans'.

Outputs (results/):
  exp2_{split}_{src}_to_{tgt}_gtrans.csv         per pair
  exp2_{split}_pairs_all_gtrans.csv              combined per split
  exp2_{split}_pair_summary_gtrans.{csv,json}    heatmap-ready aggregate
  exp2_{split}_pair_by_level_gtrans.csv          per (src, tgt, cefr_orig)

Usage:
    python exp2_gtrans_forward.py
    python exp2_gtrans_forward.py --split sentence
    python exp2_gtrans_forward.py --xlsx gtrans_templates/exp2_gtrans_input.xlsx
"""

import argparse
import json
import os
import time

import pandas as pd
import torch
from openpyxl import load_workbook
from transformers import AutoTokenizer, AutoModelForSequenceClassification

HF_TOKEN = os.environ.get("HF_TOKEN", "")
CLASSIFIER_MODEL = "UniversalCEFR/xlm-roberta-base-cefr-all-classifier"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(BASE_DIR, "gtrans_templates", "exp2_gtrans_input.xlsx")
RESULTS_DIR = os.path.join(BASE_DIR, "results")
os.makedirs(RESULTS_DIR, exist_ok=True)

SLUG = "gtrans"
CEFR_LEVELS = ["A1", "A2", "B1", "B2", "C1", "C2"]
CEFR_TO_INT = {l: i for i, l in enumerate(CEFR_LEVELS)}
INT_TO_CEFR = {i: l for i, l in enumerate(CEFR_LEVELS)}


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", default=DEFAULT_XLSX, help="Path to filled Google Translate xlsx")
    p.add_argument("--split", nargs="+", default=["sentence", "document"],
                   choices=["sentence", "document"])
    p.add_argument("--hf-token", default=None)
    p.add_argument("--limit", type=int, default=None, help="Process only first N rows (for testing)")
    return p.parse_args()


def load_xlsx_rows(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["translations"]

    # Find header row: the one whose first cell is 'row_id'
    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), start=1):
        if row and row[0] == "row_id":
            header_row_idx = i
            break
    if header_row_idx is None:
        raise RuntimeError("Could not locate header row starting with 'row_id'")

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        rid, split, src, tgt, cefr, text_orig, text_trans = row[:7]
        rows.append({
            "row_id": int(rid) if rid is not None and str(rid).strip() != "" else None,
            "split": split,
            "source_lang": src,
            "target_lang": tgt,
            "cefr_orig": cefr,
            "text_orig": text_orig,
            "text_translated": text_trans,
        })
    return rows


def load_classifier(hf_token=None):
    token = hf_token or HF_TOKEN
    os.environ["HUGGING_FACE_HUB_TOKEN"] = token
    os.environ["HF_TOKEN"] = token
    print("Loading classifier: %s" % CLASSIFIER_MODEL)
    tokenizer = AutoTokenizer.from_pretrained(CLASSIFIER_MODEL, use_auth_token=token)
    model = AutoModelForSequenceClassification.from_pretrained(CLASSIFIER_MODEL, use_auth_token=token)
    device = "cuda" if torch.cuda.is_available() else "cpu"
    model.to(device)
    model.eval()
    print("Classifier device: %s" % device)
    return tokenizer, model, device


def predict_cefr(tokenizer, model, device, text):
    if text is None or (isinstance(text, float) and pd.isna(text)) or str(text).strip() == "":
        return None
    inputs = tokenizer(str(text), return_tensors="pt", truncation=True, max_length=512)
    inputs = {k: v.to(device) for k, v in inputs.items()}
    with torch.no_grad():
        logits = model(**inputs).logits
    pred_idx = int(torch.argmax(logits, dim=-1).item())
    return model.config.id2label.get(pred_idx, INT_TO_CEFR.get(pred_idx, str(pred_idx)))


def classify_rows(rows, tokenizer, model, device):
    total = len(rows)
    t0 = time.time()
    last_log_t = t0
    for i, r in enumerate(rows):
        r["cefr_pred"] = predict_cefr(tokenizer, model, device, r["text_translated"])
        if (i + 1) % 100 == 0 or (i + 1) == total:
            now = time.time()
            elapsed = now - t0
            rate = (i + 1) / max(elapsed, 1e-6)
            eta_s = (total - (i + 1)) / max(rate, 1e-6)
            print("  classified %d/%d  (%.1f/s, elapsed %.1fs, ETA %.1fs)" % (
                i + 1, total, rate, elapsed, eta_s), flush=True)
            last_log_t = now
    return rows


def compute_delta(rows):
    for r in rows:
        orig, pred = r.get("cefr_orig"), r.get("cefr_pred")
        if orig in CEFR_TO_INT and pred in CEFR_TO_INT:
            r["delta_l"] = CEFR_TO_INT[pred] - CEFR_TO_INT[orig]
        else:
            r["delta_l"] = None
    return rows


def save_per_pair(df_split, split):
    out_paths = []
    for (src, tgt), grp in df_split.groupby(["source_lang", "target_lang"]):
        out_path = os.path.join(RESULTS_DIR, "exp2_%s_%s_to_%s_%s.csv" % (split, src, tgt, SLUG))
        out = grp.copy()
        out["direction"] = "%s->%s" % (src.upper(), tgt.upper())
        out["mt_model"] = "google_translate"
        out["source_name"] = out.get("source_name", "")
        out["format"] = split
        cols = ["direction", "source_lang", "target_lang", "source_name", "format",
                "mt_model", "cefr_orig", "text_orig", "text_translated", "cefr_pred", "delta_l"]
        out = out[[c for c in cols if c in out.columns]]
        out.to_csv(out_path, index=False, encoding="utf-8")
        out_paths.append(out_path)
    print("  wrote %d per-pair CSVs for %s-level" % (len(out_paths), split))
    return out_paths


def save_combined(df_split, split):
    out_path = os.path.join(RESULTS_DIR, "exp2_%s_pairs_all_%s.csv" % (split, SLUG))
    df = df_split.copy()
    df["direction"] = df["source_lang"].str.upper() + "->" + df["target_lang"].str.upper()
    df["mt_model"] = "google_translate"
    df["format"] = split
    cols = ["direction", "source_lang", "target_lang", "format",
            "mt_model", "cefr_orig", "text_orig", "text_translated", "cefr_pred", "delta_l"]
    df = df[[c for c in cols if c in df.columns]]
    df.to_csv(out_path, index=False, encoding="utf-8")
    print("  wrote combined %s (%d rows)" % (out_path, len(df)))
    return df


def write_pair_summary(df_all, split):
    valid = df_all.dropna(subset=["delta_l"]).copy()
    if valid.empty:
        return
    agg = (valid.groupby(["source_lang", "target_lang"])["delta_l"]
                .agg(n="count", mean="mean", std="std", median="median")
                .reset_index())
    agg["pct_up"] = (valid.assign(u=(valid["delta_l"] > 0).astype(int))
                          .groupby(["source_lang", "target_lang"])["u"].mean().values)
    agg["pct_down"] = (valid.assign(d=(valid["delta_l"] < 0).astype(int))
                            .groupby(["source_lang", "target_lang"])["d"].mean().values)
    agg["pct_preserved"] = (valid.assign(e=(valid["delta_l"] == 0).astype(int))
                                 .groupby(["source_lang", "target_lang"])["e"].mean().values)

    csv_path = os.path.join(RESULTS_DIR, "exp2_%s_pair_summary_%s.csv" % (split, SLUG))
    agg.to_csv(csv_path, index=False, encoding="utf-8")

    by_level = (valid.groupby(["source_lang", "target_lang", "cefr_orig"])["delta_l"]
                     .agg(n="count", mean="mean", std="std").reset_index())
    by_level_path = os.path.join(RESULTS_DIR, "exp2_%s_pair_by_level_%s.csv" % (split, SLUG))
    by_level.to_csv(by_level_path, index=False, encoding="utf-8")

    nested = {}
    for _, row in agg.iterrows():
        nested.setdefault(row["source_lang"], {})[row["target_lang"]] = {
            "n": int(row["n"]),
            "mean_delta_l": float(row["mean"]),
            "std_delta_l": float(row["std"]) if pd.notna(row["std"]) else None,
            "median_delta_l": float(row["median"]),
            "pct_up": float(row["pct_up"]),
            "pct_down": float(row["pct_down"]),
            "pct_preserved": float(row["pct_preserved"]),
        }
    json_path = os.path.join(RESULTS_DIR, "exp2_%s_pair_summary_%s.json" % (split, SLUG))
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(nested, f, indent=2)

    print("  wrote summary: %s" % csv_path)
    print("  wrote by-level: %s" % by_level_path)
    print("  wrote JSON:     %s" % json_path)


def print_split_summary(df_all, split):
    valid = df_all.dropna(subset=["delta_l"])
    print("\n=== %s-level summary (Google Translate) ===" % split.upper())
    print("valid rows: %d / %d" % (len(valid), len(df_all)))
    if len(valid):
        print("overall mean delta_l: %.3f" % valid["delta_l"].mean())
        print("\nmean delta_l by pair:")
        print(valid.groupby(["source_lang", "target_lang"])["delta_l"]
                   .agg(["mean", "std", "count"]).round(3).to_string())


def main():
    args = parse_args()
    print("Reading %s" % args.xlsx)
    rows = load_xlsx_rows(args.xlsx)
    if args.limit:
        rows = rows[:args.limit]
    print("Loaded %d rows" % len(rows))

    tokenizer, model, device = load_classifier(hf_token=args.hf_token)

    rows = classify_rows(rows, tokenizer, model, device)
    rows = compute_delta(rows)

    df = pd.DataFrame(rows)
    for split in args.split:
        df_split = df[df["split"] == split].copy()
        if df_split.empty:
            print("No rows for split=%s" % split)
            continue
        save_per_pair(df_split, split)
        df_all = save_combined(df_split, split)
        write_pair_summary(df_all, split)
        print_split_summary(df_all, split)


if __name__ == "__main__":
    main()
